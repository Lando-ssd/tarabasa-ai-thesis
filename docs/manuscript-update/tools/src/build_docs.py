import json
import os
import shutil
import sys

from docx_kit import *
from docx_kit import _cell_text, _shade
from docx_kit import wp as _wp
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from content_data import *
from model import MODULES, STATUS_WORD
from model_notes import FN_NOTES
import text_ab
import text_c1
import text_c2
import text_c3

OUT = "../out/pack"
os.makedirs(OUT, exist_ok=True)

# ------------------------------------------------------------------ page number checks (v29 against itself)
toc = json.load(open("../toc_check.json")) if os.path.exists("../toc_check.json") else json.load(open("../../toc_check.json"))
lotlof = json.load(open("../lot_lof_check.json")) if os.path.exists("../lot_lof_check.json") else json.load(open("../../lot_lof_check.json"))
fixed = {"Storyboard": 99, "CHAPTER I": 1}
toc = [[t, p, fixed.get(t, a)] for t, p, a in toc]
toc_bad = sum(1 for t, p, a in toc if p != a)
lot = lotlof["tables"]
lof = lotlof["figs"]
lot_bad = sum(1 for n, name, p, a in lot if p != a)
lof_bad = sum(1 for n, name, p, a in lof if p != a)
STATS = dict(toc_bad=toc_bad, toc_n=len(toc), lot_bad=lot_bad, lot_n=len(lot), lof_bad=lof_bad, lof_n=len(lof))
print(STATS)

TYPE_FILL = lambda v: TAG_FILL.get(v.split()[0])
import re as _re
_old_ids = [r[0] for r in REGISTER]
ID_MAP = {old: "R%02d" % (i + 1) for i, old in enumerate(_old_ids)}


def renum(text):
    """Register ids were numbered with gaps while the list was written; show them as R01 to R64 without gaps."""
    return _re.sub(r"R(\d\d)", lambda m: ID_MAP.get("R" + m.group(1), m.group(0)), text)


PRIO_FILL = lambda v: {"HIGH": "F9D6D3", "MEDIUM": "FDEBC8", "LOW": "E5E7EB"}.get(v)
STATUS_FILL = lambda v: {"NEW": "D9F2E3", "UPDATED": "FFF1CC", "UNCHANGED": "E8F1FB"}.get(v)


# ================================================================== 02 Insert ready text
def build_inserts():
    d = manuscript_doc()

    def wp(d, text, **kw):
        p = _wp(d, text, **kw)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.15
        return p

    h = wp(d, "TaraBasa AI: insert-ready manuscript text", bold=True, size=20, color="0B5CAD")
    wp(d, "For the writer of the manuscript. Everything below is the text to paste, in the manuscript's own formatting. Written 20 October 2026 from the system as it is today, checked against version 29 (the PDF).", size=11)
    wp(d, "How to use this file", bold=True, size=13, color="1F2937")
    wbullets(d, [
        ("Each block starts with a small grey box.", "The box is for you, not for the manuscript. It says what the block is (NEW, UPDATED, FIX, DECISION), where it goes, what it replaces, and why. Never paste the box."),
        ("Paste only what is between the box and the line 'end of ...'.", "The text, tables and pictures there already use the manuscript's formatting: Times New Roman 11 pt, justified, 1.5 line spacing, first-line indent 0.5 inch, headings in the manuscript's styles."),
        ("When you paste into the manuscript,", "choose 'Use Destination Styles' (Keep Text Only also works for text). Then nothing changes font, size or spacing."),
        ("Bracketed text [like this] is for you to complete", "(school name, results of the user test, the phone you tested on). Do not leave a bracket in the final file."),
        ("Words in italic under a table or in a grey 'Note' are instructions.", "They are not part of the manuscript."),
        ("Figure and table numbers follow the numbering map", "in 01_Review_and_Change_Register.docx (Annex 1). Change them at the end with Find and Replace."),
        ("The pictures are 300 dpi PNG files, 6.0 inches wide at most (the text width).", "The editable versions (.drawio) and vector versions (.svg) are in 05_Figures."),
    ])
    wp(d, "Tags", bold=True, size=11, color="1F2937")
    tt = d.add_table(rows=1, cols=6)
    tt.autofit = False
    for i, (tag, meaning) in enumerate([("NEW", "does not exist in v29"), ("UPDATED", "exists, replace"), ("FIX", "error to correct"), ("DECISION", "team decides first"), ("CHECK", "confirm a fact"), ("DELETE", "remove")]):
        c = tt.rows[0].cells[i]
        _cell_text(c, f"{tag}\n{meaning}", 8.5, bold=False, align="center", font=SANS)
        _shade(c, TAG_FILL[tag])
    page_break(d)
    text_ab.part_a(d)
    page_break(d)
    text_ab.part_b(d)
    page_break(d)
    text_c1.part_c1(d)
    text_c1.part_c4(d)
    text_c1.part_c5(d)
    page_break(d)
    text_c2.part_c7(d)
    page_break(d)
    text_c2.part_c8(d)
    page_break(d)
    text_c2.part_c9(d)
    page_break(d)
    text_c2.part_c10(d)
    text_c2.part_c11(d)
    page_break(d)
    text_c3.part_c12(d)
    page_break(d)
    text_c3.part_c13(d)
    text_c3.part_c14(d)
    page_break(d)
    text_c3.part_c15(d)
    text_c3.part_c16(d)
    text_c3.part_c17(d)
    page_break(d)
    text_c3.part_c18(d)
    page_break(d)
    text_c3.part_c19(d)
    text_c3.part_c20(d)
    page_break(d)
    h1(d, "PART D. BACK MATTER")
    meta_box(d, "D1", "FIX", "References, Table of Contents, List of Tables, List of Figures, Appendices list.",
             "See 04_New_References_and_Basis.docx for the references and Annexes 1 and 2 of 01 for the numbers.",
             note="Do this last. The contents, the list of tables and the list of figures are typed tables in version 29, so their page numbers do not follow the text.")
    wp(d, "1. References: use 04. 2. Renumber figures and tables with the map. 3. Rebuild the contents (best from the heading styles). 4. Fix the page numbers in the two lists. 5. Read the whole file once for the words SMS, Kotlin, Bootstrap, XAMPP, Cloud SQL, Google Cloud, Semaphore, Android Studio: none should remain except where they say the system does not use them.", size=10.5)
    end_marker(d, "D1")
    d.save(OUT + "/02_Insert_Ready_Text.docx")


# ================================================================== 01 Register
def build_register():
    d = work_doc(landscape=True, margins=0.7)
    d.add_heading("Review and Change Register", level=1)
    wp(d, "Manuscript TARABASA-AI-Final-Manuscript29.pdf (188 pages; pages below are the numbers printed on the pages). Reviewed 20 October 2026 against the system as built. Every row says what is wrong or missing, what to do, and which block of 02_Insert_Ready_Text.docx has the replacement text.")
    from collections import Counter
    types = Counter(r[4] for r in REGISTER)
    prios = Counter(r[5] for r in REGISTER)
    wp(d, f"{len(REGISTER)} findings. By type: " + ", ".join(f"{k} {v}" for k, v in sorted(types.items())) + ". By priority: " + ", ".join(f"{k} {prios[k]}" for k in ("HIGH", "MEDIUM", "LOW")) + ".", bold=True)
    d.add_heading("What matters most (the 9 things a panel will find first)", level=2)
    wbullets(d, [
        ("SMS and Google Speech-to-Text are everywhere", "(business model, sprint 9, Gantt, network, technology stack, software list, comparative matrix). Neither is in the system: it uses Vosk through a reading service and sends emails only."),
        ("The technology stack is the old plan, not the build:", "Kotlin, XAMPP, Bootstrap and Cloud SQL against React Native with Expo, Laravel 13, TiDB Cloud on Render."),
        ("The workflow figures show a learner moved up or down by one reading,", "a Parent starting a reading, a Teacher browsing the Open Repository and a Pending teacher blocked at sign in. None of that is how it works."),
        ("The rule-based adaptive support is not explained anywhere,", "although it is the core idea. New section C9."),
        ("The FDD has nine modules and misses the adaptive rules and the Admin screens.", "A new eleven-module FDD is ready for the panel (03)."),
        ("The database figures show 16 tables; the database has 23.", "'Curriculum Guide' and 'Badge' are described as tables but do not exist."),
        ("Chapter 1 'Flow of the Study' still has a sentence about hand gestures", "from another study."),
        ("No testing evidence and no data privacy section", "in a thesis about a system that holds children's data."),
        ("Three sources are cited but missing from the references,", "and the contents, list of tables and list of figures have wrong page numbers for most entries."),
    ])
    d.add_heading("Register", level=2)
    rows = []
    for r in REGISTER:
        rid, where, finding, action, typ, prio, ins = r
        finding = finding.format(**STATS)
        rows.append([ID_MAP[rid], where, renum(finding), renum(action), typ, prio, ins])
    wtable(d, ["ID", "Where in v29", "What is wrong or missing", "What to do", "Type", "Priority", "Block"], rows, [0.5, 1.7, 3.6, 2.2, 0.8, 0.7, 0.6],
           size=8.5, fills=[None, None, None, None, TYPE_FILL, PRIO_FILL, None], total_w=9.6)
    d.add_heading("Checked and left alone", level=2)
    wp(d, renum("Read and found no conflict with the system: Acknowledgement and Dedication; Chapter 1 introduction, rationale, objectives (apart from R15) and significance; Chapter 2 theoretical background, related literature and related studies (apart from R23, R30, R34 and R35); Agile sprint plan tables 3.1 to 3.4 and the columns of sprints 6 to 8 and 10 (apart from SMS words in 9); the Validation Board experiments (apart from R49); Gantt dates; the Parent registration screen design (apart from the name fields). This is a reading check, not a line by line proof: the writer should still read each page while pasting."))
    page_break(d)
    d.add_heading("Annex 1. Numbering map (old number in version 29, new number after the inserts)", level=1)
    wp(d, "Use Word's Find and Replace from the highest number down so that no number is changed twice.", italic=True)
    wtable(d, ["Figure in v29", "Figure after", "What it is"], [list(x) for x in FIG_MAP], [1.6, 1.6, 6.4], size=9, total_w=9.6)
    page_break(d)
    wtable(d, ["Table in v29", "Table after", "What it is"], [list(x) for x in TABLE_MAP], [1.8, 1.8, 6.0], size=9, total_w=9.6)
    page_break(d)
    d.add_heading("Annex 2. Page numbers in the contents and lists, as printed in v29 and where the heading really is", level=1)
    wp(d, "These change again after the inserts. They show the size of the problem and what to fix last. 'Real' is the page number printed on the page where the heading or caption sits in version 29.", italic=True)
    wtable(d, ["Contents entry", "Says", "Real"], [[t, str(p), str(a)] for t, p, a in toc], [4.0, 0.8, 0.8], size=8.5, fills=[None, None, lambda v: None], total_w=5.6)
    d.add_paragraph()
    wtable(d, ["List of Tables entry", "Says", "Real"], [[f"Table {n}: {name}", str(p), str(a)] for n, name, p, a in lot], [4.0, 0.8, 0.8], size=8.5, total_w=5.6)
    d.add_paragraph()
    wtable(d, ["List of Figures entry", "Says", "Real"], [[f"Figure {n}: {name}", str(p), str(a)] for n, name, p, a in lof], [4.0, 0.8, 0.8], size=8.5, total_w=5.6)
    d.save(OUT + "/01_Review_and_Change_Register.docx")


# ================================================================== 03 FDD panel review
PANEL_QUESTIONS = [
    ("1", "'AI' in TaraBasa means a speech recognition service plus written rules and a formula service. No model was trained. Is that acceptable for the title 'AI-powered adaptive reading intervention'?", "Chapter 1 scope; Figure 9; Table 7"),
    ("2", "One class holds one grade only, even for a teacher who handles several grades (a DepEd multigrade class is not modelled).", "Figure 4.2; Module 3"),
    ("3", "A reading step moves only after 3 strong readings (90 percent or more) of at least 2 activities, or after 2 weak readings (under 70) in a row, and only one step.", "Figure 3.3; Table 7"),
    ("4", "A Parent cannot start a child's reading. A child reads only after signing in with the PIN.", "Figure 5.2; Module 6"),
    ("5", "A teacher who is not yet approved can sign in with limited access (view only, 2 free AI drafts) after confirming the email.", "Figure 4.1; Module 1"),
    ("6", "Only Parents browse the Open Repository; Teachers only share. The payment is simulated; no real money is handled.", "Figure 4.3, 5.2; Module 5"),
    ("7", "The Admin works on the website only and has Overview, Approvals, Accounts, System health and Activity log. There is no AI configuration screen, no announcements and no admin reports.", "Figure 6.1; Module 11"),
    ("8", "There is no SMS. Parents and teachers see updates in the app and the website. The digital equity goal (parents without smartphones) becomes future work.", "Chapter 3 network; Module 9"),
    ("9", "The first reading check starts where the Parent's profile says the child is, not at the grade, and has at most 3 items.", "Figure 3.2; Table 7"),
    ("10", "Practice Games are free play: they never change points, streak or reading step.", "Figure 3.4; Module 7"),
    ("11", "Linking an existing child needs the child's PIN as well as the code (the original actor prompt asked for the code only).", "Figure 5.1; Module 2"),
    ("12", "A parent or teacher can delete their own account, and a parent can delete a child, without anyone's approval (Data Privacy Act).", "Figure 5.1; Module 1, 2, 11"),
]

WORKFLOW_REVIEW = [
    ["3.1 and 3.2 Learner", "Login by PIN only. One reading moves the level up, keeps it, or moves it down with a flag. 'Warm-Up Game'. No comprehension quiz, word bank, badges detail, or first-check ladder. 'AI Score (Speech-to-Text)'.", "3.2 to 3.4 (plus the overview 3.1): card or code with PIN; first reading check ladder; practice stage; unclear audio retry; comprehension quiz; slow step rule; word bank; recommender; badges; alerts; games; bookshelf; rewards."],
    ["4.1 and 4.2 Teacher", "A not yet approved teacher is stopped with a pending notice. The Teacher can browse the Open Repository. No alerts, no analytics detail, no promotion flow, no reading groups, no fit check.", "4.1 to 4.4: confirm email; Pending can view; grade requests; one grade per class; add by code or QR with a check; roster, move, promotion; generate in the background; place in levels; assign with a fit check; share; analytics; alerts; profile."],
    ["5.1 and 5.2 Parents", "A 'Practice Reading' branch opens the learner interface. No email confirmation, no PIN reset, no QR card, no removing a child, no alerts detail.", "5.1 and 5.2: confirm email; add a child with the 3 questions; learner card; link with PIN; reset PIN; remove a child; progress with tips; activities; alerts; delete account."],
    ["6.1 Administrator", "Edit Profile and Save Changes branches (no such page). Only approve or reject teachers and activate or deactivate accounts.", "6.1: Approvals (waiting, rejected, grade requests); Accounts (search, deactivate, confirm email, delete); System health; Activity log."],
]


def build_panel():
    d = work_doc(landscape=True, margins=0.7)
    d.add_heading("FDD and Use Case: panel review sheet", level=1)
    wp(d, "Version 2 of the Functional Decomposition Diagram (FDD) for the panel to check against how TaraBasa AI really works. Written 20 October 2026. The picture below is the review version: green is new, amber is updated, blue is unchanged from version 29. The manuscript version of each figure is one colour.")
    from docx_kit import picture
    picture(d, "../out/fdd/FDD_00_overview_review.png", 7.0, 5.0)
    from collections import Counter
    tot = Counter()
    for m in MODULES:
        for _a, fns in m["actors"]:
            for _n, s in fns:
                tot[s] += 1
    wp(d, f"Version 29 has 9 modules and 80 functions. Version 2 has 11 modules and {sum(tot.values())} functions: {tot['new']} new, {tot['upd']} updated, {tot['same']} unchanged. Seven functions of version 29 were merged into others (listed below). The use case diagrams are drawn from the same list, so each use case is a function of the FDD.", bold=True)
    d.add_heading("Questions for the panel", level=2)
    wp(d, "Each line is something the system does today that differs from the original plan or from the actor prompts. Please mark each one.")
    rows = [[n, q, where, "", "", ""] for n, q, where in PANEL_QUESTIONS]
    wtable(d, ["No.", "Statement about the flow", "Shown in", "Agree", "Change", "Comment"], rows, [0.4, 5.0, 1.6, 0.5, 0.6, 1.5], size=9, total_w=9.6)
    d.add_paragraph()
    wp(d, "Order of the modules. In version 29 the FDD figure lists Reading Assessment fifth, Gamified Learning sixth and Open Repository seventh, but the List of Modules numbers Open Repository 5, Reading Assessment 6 and Gamified Learning 7. Version 2 uses the List of Modules numbers in both.", italic=True)
    page_break(d)
    d.add_heading("From version 29 to version 2", level=2)
    wtable(d, ["Version 29", "Version 2", "What changed"],
           [["Module 1 Account Management (27 functions)", "Module 1 (26 functions)", "Email must be confirmed; card or code and PIN for learners; Google sign in; grade change request; self deletion. The Admin's 6 account functions moved to Module 11."],
            ["Module 2 Learner Management (6)", "Module 2 (10)", "Check-digit Learner Code, QR card, PIN reset, removing a child, profile details."],
            ["Module 3 Class Management (7)", "Module 3 (9)", "One grade per class, QR add with a check, reading groups, learner profile, moving a learner."],
            ["Module 4 Reading Activity (6)", "Module 4 (7)", "Levels per request, background writing, writing for a class, placing in levels, fit check."],
            ["Module 5 Open Repository (6)", "Module 5 (6)", "Best match first; payment simulated."],
            ["Module 6 Learner Reading Assessment (10)", "Module 6 (9)", "Placement and diagnostic became one first reading check; warm-up and practice added; result and feedback merged."],
            ["Module 7 Gamified Learning (6)", "Module 7 (7)", "100 badges, day streak, four games, text size."],
            ["Module 8 Dashboard Analytics (8)", "Module 8 (10)", "'Did it work?', class report, plain-language progress."],
            ["Module 9 Notification Management (4)", "Module 9 (8)", "Alerts from readings, one-tap suggested activity, handled, emails (no SMS)."],
            ["(none)", "Module 10 Adaptive Reading Support (10, all new)", "The rules and services the system applies to each reading."],
            ["(Admin part of Module 1)", "Module 11 System Administration and Privacy (10)", "Approvals, grade requests, accounts, deletion, system health, activity log."]],
           [2.6, 2.4, 4.6], size=9, total_w=9.6)
    d.add_paragraph()
    wp(d, "Functions of version 29 that were merged or renamed:", bold=True)
    wtable(d, ["In version 29", "In version 2"],
           [["Manage user accounts; Deactivate account", "Deactivate or reactivate account (Module 11)"],
            ["View pending teacher accounts; Verify employee ID; Activate teacher account; Reject teacher account", "Approve or reject teacher (with reason) (Module 11)"],
            ["Take placement assessment; Take diagnostic assessment", "Take first reading check (Module 6)"],
            ["Start adaptive recommended activity; Start reading activity", "Start recommended or other activity (Module 6)"],
            ["View reading feedback; View assessment result", "View result and word review (Module 6)"],
            ["View adaptive learning progress", "View reading path and progress (Module 6)"],
            ["Verify email", "Confirm email by link (Module 1)"]],
           [5.0, 4.6], size=9, total_w=9.6)
    d.add_paragraph()
    page_break(d)
    d.add_heading("Module by module (every function, with its status)", level=2)
    for m in MODULES:
        d.add_heading(f"Module {m['no']}: {m['name']}  ({m['change']})", level=3)
        wp(d, m["why"], italic=True)
        fn = f"../out/fdd/FDD_{m['no']:02d}_{m['name'].lower().replace(' ', '_')}_review.png"
        picture(d, fn, 7.2, 4.4)
        rows = []
        for actor, fns in m["actors"]:
            for name, st in fns:
                rows.append([actor, name, STATUS_WORD[st], FN_NOTES.get(name, "")])
        wtable(d, ["Actor", "Function", "Status", "What is different from version 29"], rows, [0.9, 2.6, 0.9, 5.2], size=8.5,
               fills=[None, None, STATUS_FILL, None], total_w=9.6)
        d.add_paragraph()
    page_break(d)
    d.add_heading("Workflow figures of version 29 that were wrong", level=2)
    wp(d, "Version 29's workflow figures were read as pictures against the system's routes. The new figures are in 05_Figures and in block C4 of the insert text.")
    wtable(d, ["Old figure", "What was wrong", "New figures"], WORKFLOW_REVIEW, [1.4, 4.4, 3.8], size=9, total_w=9.6)
    d.add_paragraph()
    d.add_heading("Use case diagrams", level=2)
    wp(d, "The six use case diagrams are drawn from the same list as the FDD. Colour in the review versions is the same as above. Version 29 had four diagrams with 8 + 8 + 21 + 15 use cases; version 2 has six diagrams. The same function used by two actors (for example Register account) is one use case with two lines.")
    d.save(OUT + "/03_FDD_and_Use_Case_Panel_Review.docx")


# ================================================================== 04 References
BASIS = [
    ["The speech engine is Vosk, built on Kaldi", "Povey et al. (2011)", "RECALLED"],
    ["Children's speech is harder to recognize, so scores are cleaned and not trusted blindly", "Potamianos & Narayanan (2003)", "RECALLED"],
    ["The speed score is not grade adjusted (stated limit)", "Hasbrouck & Tindal (2017) (as the thing the score is NOT based on)", "LISTED"],
    ["Activity drafts come from a generative model that a teacher must approve", "Gemini Team (2023)", "LISTED"],
    ["Tap-only controls for ages 7 to 8", "Yadav et al. (2020)", "LISTED"],
    ["Games make the child read aloud; tapping alone is not enough", "Hirsh-Pasek et al. (2015)", "LISTED"],
    ["Syllable method: split, join, read in a sentence", "Liberman et al. (1974); Therrien (2004); Department of Education (2024a)", "RECALLED"],
    ["Phil-IRI levels (Independent, Instructional, Frustration) by oral reading accuracy", "Department of Education (2018)", "RECALLED"],
    ["Right to have personal data removed", "Republic Act No. 10173 (2012)", "RECALLED"],
    ["Learning analytics dashboards: effect on participation, small effect on achievement (already cited)", "Kaliisa et al. (2024)", "VERIFIED"],
    ["Parent involvement and reading achievement (already cited)", "Hernandez & Decena (2024)", "VERIFIED"],
]
STATUS_REF_FILL = lambda v: {"VERIFIED": "D9F2E3", "LISTED": "FFF1CC", "RECALLED": "FADBD8"}.get(v)


def build_refs():
    d = work_doc(landscape=False, margins=0.8)
    d.add_heading("New references, corrections and the basis for each choice", level=1)
    wp(d, "This file is kept apart from the manuscript on purpose. Part 2 is the only part to paste into the reference list; everything else is for the writer. Written 20 October 2026.")
    d.add_heading("1. Check status: read this before you trust a source", level=2)
    wtable(d, ["Status", "Meaning"],
           [["VERIFIED", "The record itself (journal page, publisher or index record) was opened in this review and the facts below were read there."],
            ["LISTED", "Found in a search listing; authors, journal and year agree across listings. The abstract and the exact sentence were not read: open the paper before keeping the sentence that cites it."],
            ["RECALLED", "From the team's earlier notes or from memory; at most a search listing. Confirm every detail (title, DOI, pages) at the publisher before printing."]],
           [1.0, 6.0], size=9, fills=[STATUS_REF_FILL, None], total_w=6.9)
    d.add_heading("2. Reference entries to add (paste into the reference list in alphabetical order, APA 7, hanging indent 0.5 inch)", level=2)
    from docx.shared import Inches as In
    for rid, apa, status, used, note in sorted(NEW_REFS, key=lambda x: x[1]):
        p = d.add_paragraph()
        p.paragraph_format.left_indent = In(0.5)
        p.paragraph_format.first_line_indent = In(-0.5)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text_ab.smart(apa))
        from docx_kit import _set_font
        _set_font(r, TNR, 11)
    wp(d, "(The reference list in the manuscript is Times New Roman 11 pt, 1.5 line spacing, hanging indent. Pasting with 'Use Destination Styles' keeps that.)", italic=True)
    d.add_heading("3. Where each new entry is used, and how far it was checked", level=2)
    rows = [[rid, apa.split("(")[0].strip()[:42] + " (" + apa.split("(")[1].split(")")[0] + ")", status, used, note] for rid, apa, status, used, note in NEW_REFS]
    wtable(d, ["No.", "Source", "Status", "Used in", "What to check"], rows, [0.4, 1.5, 0.8, 1.6, 2.8], size=8.5, fills=[None, None, STATUS_REF_FILL, None, None], total_w=6.9)
    d.add_heading("4. Corrections to citations and entries already in version 29", level=2)
    wtable(d, ["Where", "Problem", "Do this"], [list(x) for x in REF_FIXES], [1.6, 2.6, 2.7], size=8.5, total_w=6.9)
    d.add_heading("5. Which source supports which feature", level=2)
    wp(d, "A single table for the defense: a feature, the source behind it, and how far the source was checked. A source marked RECALLED must be read before the table is shown to a panel.")
    wtable(d, ["Feature or claim", "Source", "Status"], BASIS, [3.6, 2.6, 0.9], size=8.5, fills=[None, None, STATUS_REF_FILL], total_w=6.9)
    d.add_heading("6. Optional additional basis", level=2)
    wtable(d, ["No.", "Reference", "Status", "Support for"], [[a, b, c, e] for a, b, c, e in OPTIONAL_REFS], [0.4, 3.6, 0.8, 2.1], size=8.5, fills=[None, None, STATUS_REF_FILL, None], total_w=6.9)
    d.add_heading("7. What was not found", level=2)
    wbullets(d, [
        "No study was found that shows closed syllables must be taught before open syllables. The system's choice (closed syllables first, no open syllables such as 'baby') is common practice in structured literacy programs; say it is the team's choice, not a finding.",
        "No source was found for the numbers in Table 7 (90, 70, 3, 2, 0.50, 25 percent, the word minimums). They are the team's own settings. The text says so; keep it that way.",
        "No public recording set of Filipino children's speech was found. Public children's corpora (CMU Kids, MyST, OGI Kids) are English only. None is used.",
        "DepEd's own pages did not open from the build machine (403). The DepEd order numbers and titles in the project notes came from search results and must be copied from deped.gov.ph.",
    ])
    d.save(OUT + "/04_New_References_and_Basis.docx")


# ================================================================== 00 README
FIGURE_INDEX = [
    ["Figure 3.1", "F01_system_flow", "System Flow Chart"],
    ["Figure 3.2 to 3.4", "F05_learner_workflow_1, F06_learner_workflow_2, F07_learner_workflow_3", "Program Workflow (Learners)"],
    ["Figure 4.1 to 4.4", "F08_teacher_workflow_1 to F11_teacher_workflow_4", "Program Workflow (Teacher)"],
    ["Figure 5.1, 5.2", "F12_parent_workflow_1, F13_parent_workflow_2", "Program Workflow (Parents)"],
    ["Figure 6.1", "F14_admin_workflow", "Program Workflow (Administrator)"],
    ["Figure 7.1 to 7.12", "FDD_00_overview_clean, FDD_01 to FDD_11 (_clean)", "Functional Decomposition Diagram (manuscript version; the _review versions show new, updated and unchanged)"],
    ["Figure 8.1 to 8.6", "UC_1_account to UC_6_adaptive", "Use Case Diagrams"],
    ["Figure 9", "F15_adaptive_rule_chain", "Adaptive Reading Support Rule Chain"],
    ["Figure 42", "F16_erd_version2", "Entity Relationship Diagram"],
    ["Figure 43", "F03_network_model", "Network Model"],
    ["Figure 44", "F04_network_topology", "Network Topology"],
    ["Figure 45", "F02_architecture_tech_stack", "System Architecture and Technology Stack"],
]
SPEC = [
    ["Paper", "US Letter, 8.5 x 11 inch, portrait"],
    ["Margins", "Left 1.5 inch, right 1 inch, top 1 inch, bottom 1 inch (text width 6.0 inch)"],
    ["Body text", "Times New Roman 11 pt, justified, line spacing 1.5, no space before or after, first line indent 0.5 inch"],
    ["Heading 1 (CHAPTER)", "Times New Roman 14 pt bold, centered, line spacing 1.5"],
    ["Heading 2 (chapter title, APPENDIX)", "Times New Roman 12 pt bold, centered"],
    ["Heading 3 (section)", "Times New Roman 11 pt bold, left, 10 pt before"],
    ["Table caption", "Above the table, centered, two lines: 'Table 31', then the title in CAPITALS (style Heading 4: Times New Roman 11 pt, not bold)"],
    ["Figure caption", "Below the figure, centered: 'Figure 7.2: ' in plain type, then the title in bold and underlined"],
    ["Table", "Full width, thin black borders (0.5 to 1 pt), Times New Roman; 8.5 pt for wide tables, 11 pt for two-column tables; header row bold and repeated on each page"],
    ["Picture", "PNG at 300 dpi, centered, 6.0 inch wide at most, caption right below; keep a figure and its caption on the same page"],
    ["Hanging indent (references)", "0.5 inch, Times New Roman 11 pt"],
    ["Quotation marks", "Typographic (curly) quotes and apostrophes; no dash used as a pause inside new text (use commas or a new sentence)"],
]
DECISIONS = [
    ["Which school is the study site?", "Chapter 1 says Cebu City; the Validation Board says Suba Elementary School in Liloan.", "C5, R18"],
    ["How many respondents answered each survey?", "Needed for every percentage in the Validation Board.", "C5, R49"],
    ["Keep the old UI figures as 'initial design' or replace with real screenshots?", "Three Admin screens were never built.", "C10, R56 to R59"],
    ["One reference list, or grouped by type?", "Ask the adviser.", "04, R75"],
    ["What to do with 'No Read, No Move'?", "Cite the DepEd issuance or remove (C3 removes it).", "R45"],
    ["Is 'rule-based' acceptable wording for the adaptive part?", "The panel should see Figure 9 and Table 7.", "03, question 1"],
    ["Programmers for modules 10 and 11", "Version 29 has owners for modules 1 to 9.", "C18"],
    ["Real results of the user acceptance test and the phone test", "Bracketed lines in C19.", "C19"],
]


def build_readme():
    d = work_doc(landscape=False, margins=0.9)
    wp(d, "TaraBasa AI: manuscript update pack", bold=True, size=22, color="0B5CAD", after=2)
    wp(d, "For the writer of the manuscript and the team. Written 20 October 2026. Start here.", size=11, after=8)
    d.add_heading("What this pack is", level=2)
    wp(d, "The manuscript (version 29, 188 pages) was reviewed against the system as it is built today. This pack separates what must change into files, so the writer can work through them without searching. Nothing in the manuscript was edited; every new or changed piece is here, in the manuscript's own formatting, with the place it goes.")
    d.add_heading("The files", level=2)
    wtable(d, ["File", "What it holds", "Who uses it"],
           [["00_README_START_HERE.docx", "This page: work order, formatting rules, decisions, limits.", "Everyone"],
            ["01_Review_and_Change_Register.docx", f"{len(REGISTER)} findings with page, action, type and priority; numbering map (Annex 1); page number errors (Annex 2).", "Writer, adviser"],
            ["02_Insert_Ready_Text.docx", "All new and replacement text, tables and pictures, in manuscript formatting. Each block says where it goes. Parts A to D.", "Writer"],
            ["03_FDD_and_Use_Case_Panel_Review.docx", "The new eleven-module FDD with every function marked new, updated or unchanged; questions for the panel; wrong workflow figures.", "Panel, team"],
            ["04_New_References_and_Basis.docx", "New references (APA), how far each was checked, corrections to existing citations, feature to source table. Kept apart from the manuscript.", "Writer, adviser"],
            ["05_Figures/", "Every new figure as PNG (for Word), SVG (vector) and .drawio (editable at app.diagrams.net). Colour review versions of the FDD and use cases are in review_versions.", "Writer, team"]],
           [3.0, 3.1, 0.8], size=9, total_w=6.9)
    d.add_heading("Work order", level=2)
    for i, t in enumerate([
        "Settle the decisions below (they block several blocks).",
        "Fix the wrong facts: Part A, Part B, C1 to C3, C6 (find SMS, Google Speech-to-Text, Kotlin, XAMPP, Bootstrap, Cloud SQL).",
        "Replace the figures and their paragraphs: workflows (C4), FDD (C7), use cases (C8), ERD and tables (C12), network (C13).",
        "Add the new sections: B3, C9, C19, C20.",
        "Replace the stack, software, hardware and module tables: C15 to C18.",
        "Add the references and make the corrections (04).",
        "Renumber figures and tables with the map (Annex 1 of 01), rebuild the contents, fix the page numbers in the lists.",
        "Read the whole file once. Search for SMS, Kotlin, Bootstrap, XAMPP, Cloud SQL, Google Cloud, Semaphore, Android Studio.",
    ], start=1):
        wp(d, f"{i}. {t}", after=2)
    d.add_heading("Decisions the team must make first", level=2)
    wtable(d, ["Question", "Why", "Where"], [[a, b, renum(c)] for a, b, c in DECISIONS], [3.0, 2.8, 1.1], size=9, total_w=6.9)
    d.add_heading("Formatting rules (copied from the manuscript's Word file)", level=2)
    wp(d, "These were read from the manuscript's Word styles and checked against the PDF. The insert text already follows them. If you type new text, use them.")
    wtable(d, ["Item", "Rule"], SPEC, [1.8, 5.1], size=9, total_w=6.9)
    d.add_heading("Figures: how to use them", level=2)
    wbullets(d, [
        "In Word: Insert, Pictures, This Device, choose the PNG, then set 'Center' and 'In Line with Text'. Width is already 6.0 inch or less; do not stretch.",
        "To change a figure: open the .drawio file at app.diagrams.net (File, Open from, Device), edit the text or boxes, then File, Export as, PNG (set zoom 300 percent) or SVG. Word can also insert the SVG directly.",
        "The FDD and use case figures are generated from one list of functions. If a function changes, change it in both the FDD and the use case, and the List of Modules table.",
        "Every picture is on a white background and uses colours that print well in grey.",
    ])
    wtable(d, ["Figure number after the inserts", "File name(s) in 05_Figures", "What it is"], FIGURE_INDEX, [1.6, 3.0, 2.3], size=8.5, total_w=6.9)
    d.add_heading("What was checked and how", level=2)
    wbullets(d, [
        "The manuscript was read from the PDF text (188 pages) and its pictures (FDD, use cases, workflows, ERD) were looked at as images. The Word file in the Downloads folder is from July 2026 and older than the PDF, so the PDF is the version reviewed; its styles were used for the formatting rules.",
        "Every statement about the system was checked against the code, the routes and the database migrations in the repository on 20 October 2026, and against the project notes.",
        "Page numbers are the numbers printed on the pages of version 29.",
        "The numbers and settings in Table 7 are read from the configuration files; the text says they are the team's own choices.",
    ])
    d.add_heading("What is not done or not proven", level=2)
    wbullets(d, [
        "No real child, real classroom or real phone test has been done; Block C19 has bracketed lines for the results. The phone test is planned for 21 October 2026.",
        "Several references are marked RECALLED or LISTED: read them before the defense (04).",
        "The Validation Board numbers (n) and the study school are for the team to confirm (C5).",
        "Real screenshots for the 'Implemented Screens' section are not included (C10 lists which to take).",
        "Nothing was pushed or published. The pack is a local folder and a zip.",
    ])
    d.save(OUT + "/00_README_START_HERE.docx")


def copy_figures():
    base = OUT + "/05_Figures"
    if os.path.exists(base):
        shutil.rmtree(base)
    for sub in ("for_Word_PNG", "svg", "editable_drawio", "review_versions"):
        os.makedirs(f"{base}/{sub}", exist_ok=True)
    n = 0
    for folder in ("figures", "fdd", "usecase"):
        src = f"../out/{folder}"
        for f in sorted(os.listdir(src)):
            path = f"{src}/{f}"
            ext = f.rsplit(".", 1)[-1]
            review = "_review" in f
            if review:
                shutil.copy(path, f"{base}/review_versions/{f}")
            elif ext == "png":
                shutil.copy(path, f"{base}/for_Word_PNG/{f}")
            elif ext == "svg":
                shutil.copy(path, f"{base}/svg/{f}")
            elif ext == "drawio":
                shutil.copy(path, f"{base}/editable_drawio/{f}")
            n += 1
    print("figures copied", n)


if __name__ == "__main__":
    build_inserts()
    build_register()
    build_panel()
    build_refs()
    build_readme()
    copy_figures()
    print("done")
